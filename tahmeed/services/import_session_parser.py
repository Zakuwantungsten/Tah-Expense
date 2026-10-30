"""Read many daily / MASTER Excel workbooks into rows grouped by reconciled day.

Batch (reconciled) date comes from the source file name, e.g.
``MATUMIZI YA 19-10-26.xlsx``: first the SOURCE FILE column of a MASTER
workbook, then the workbook filename, then the majority row date of the file.
The Excel row DATE is kept on each row but never decides the batch.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence

import openpyxl

from tahmeed.models.transaction import pack_money
from tahmeed.services.daily_import_service import (
    DailyImportCancelled,
    cell_str,
    detect_date_from_name,
    parse_amount,
    parse_date_value,
)

SKIP_EMPTY_DESCRIPTION = "empty_description"
SKIP_TOTAL_ROW = "total_row"
SKIP_MISSING_AMOUNT = "missing_amount"
SKIP_UNPARSEABLE_AMOUNT = "unparseable_amount"
SKIP_ZERO_AMOUNT = "zero_amount"
SKIP_BAD_SOURCE_DATE = "bad_source_date"

SKIP_REASON_LABELS: Dict[str, str] = {
    SKIP_EMPTY_DESCRIPTION: "Empty description",
    SKIP_TOTAL_ROW: "Total row",
    SKIP_MISSING_AMOUNT: "Missing amount",
    SKIP_UNPARSEABLE_AMOUNT: "Unparseable amount",
    SKIP_ZERO_AMOUNT: "Zero amount",
    SKIP_BAD_SOURCE_DATE: "No reconciled date",
}

_HEADER_ALIASES: Dict[str, tuple[str, ...]] = {
    "serial": ("s/no", "sno", "s no", "serial", "#"),
    "date": ("date",),
    "description": ("description", "desc"),
    "truck": ("truck no", "truck number", "truckno", "truck", "class"),
    "lpo": ("lpo nos", "lpo", "lpo/do"),
    "do": ("do no", "do number"),
    "memo": ("memo",),
    "notes": ("notes", "note", "ref_float", "ref float"),
    "tzs": ("tzs",),
    "usd": ("usd", "us$", "dollar"),
    "receipt": ("receipt status", "receipt", "rcpt"),
    "ownership": ("ownership", "own"),
    "approver": ("apr by", "approver", "apr"),
    "source_file": ("source file", "sourcefile", "source"),
}

_CLASSIC_MATUMIZI_COLS: Dict[str, int] = {
    "serial": 0,
    "date": 1,
    "description": 3,
    "truck": 4,
    "lpo": 5,
    "do": 6,
    "memo": 9,
    "notes": 10,
    "tzs": 11,
    "usd": 12,
    "receipt": 13,
    "ownership": 14,
    "approver": 15,
}

_HEADER_SCAN_ROWS = 20

_MATUMIZI_DATE_RE = re.compile(
    r"MATUMIZI\s+YA\s+(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})",
    re.IGNORECASE,
)


@dataclass
class StagedRow:
    """One usable Excel row waiting in an import session."""

    row_id: str
    file_key: str
    source_file: str
    excel_row: int
    batch_date: date
    description: str
    amount: float
    currency: str = "TZS"
    amount_usd: Optional[float] = None
    excel_date: Optional[date] = None
    truck_raw: str = ""
    lpo_do: str = ""
    do_number: str = ""
    memo: str = ""
    notes: str = ""
    receipt_status: str = ""
    ownership: str = ""
    approver: str = ""
    serial: Optional[int] = None
    excel_cells: Dict[str, str] = field(default_factory=dict)


@dataclass
class SkippedExcelRow:
    file_key: str
    source_file: str
    excel_row: int
    reason: str
    detail: str = ""
    batch_date: Optional[date] = None
    excel_cells: Dict[str, str] = field(default_factory=dict)

    @property
    def reason_label(self) -> str:
        return SKIP_REASON_LABELS.get(self.reason, self.reason.replace("_", " ").title())


@dataclass
class LoadedFile:
    file_key: str
    path: str
    name: str
    sheet: str
    headers: List[str] = field(default_factory=list)
    row_count: int = 0
    skipped_count: int = 0
    warnings: List[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    file: LoadedFile
    rows: List[StagedRow] = field(default_factory=list)
    skipped: List[SkippedExcelRow] = field(default_factory=list)


def file_key_for(path: str | Path) -> str:
    return str(Path(path).expanduser().resolve()).casefold()


def normalize_import_description(value: Any) -> str:
    """Uppercase, collapse whitespace, drop trailing dashes left after a plate."""
    text = " ".join(cell_str(value).upper().split())
    return re.sub(r"[\s\-–—]+$", "", text).strip()


def parse_source_file_date(name: str) -> Optional[date]:
    """Reconciled date from ``MATUMIZI YA dd-mm-yy`` or any ``dd-mm-yyyy`` name."""
    text = str(name or "")
    match = _MATUMIZI_DATE_RE.search(text)
    if match:
        day, month, year = (int(part) for part in match.groups())
        if year < 100:
            year += 2000
        try:
            return date(year, month, day)
        except ValueError:
            return None
    return detect_date_from_name(text)


def format_cell_display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value).strip()


def _norm_header(value: Any) -> str:
    text = cell_str(value).lower().replace(".", " ")
    return " ".join(text.split())


def _map_headers(header: Sequence[Any]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for index, cell in enumerate(header):
        norm = _norm_header(cell)
        if not norm:
            continue
        spaced = " ".join(norm.replace("_", " ").replace("-", " ").split())
        for field_name, aliases in _HEADER_ALIASES.items():
            if field_name in mapping:
                continue
            if norm in aliases or spaced in aliases:
                mapping[field_name] = index
                break
    return mapping


def _looks_like_classic_matumizi(header: Sequence[Any]) -> bool:
    cells = [_norm_header(c) for c in header]
    if len(cells) < 4:
        return False
    return cells[1].startswith("date") and "desc" in cells[3]


def _is_blank_row(values: Sequence[Any]) -> bool:
    return all(v is None or str(v).strip() == "" for v in values)


def _cell(values: Sequence[Any], cols: Dict[str, int], key: str) -> Any:
    index = cols.get(key)
    if index is None or index >= len(values):
        return None
    return values[index]


def _excel_cells(header: Sequence[Any], values: Sequence[Any]) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for index, label in enumerate(header):
        name = cell_str(label)
        if not name or name in out:
            continue
        out[name] = format_cell_display(values[index] if index < len(values) else None)
    return out


def _header_labels(header: Sequence[Any]) -> List[str]:
    return list(dict.fromkeys(cell_str(c) for c in header if cell_str(c)))


def _parse_serial(raw: Any) -> Optional[int]:
    try:
        if raw is not None and cell_str(raw):
            return int(float(raw))
    except (TypeError, ValueError):
        pass
    return None


def list_sheet_names(path: str | Path) -> List[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def common_sheet_names(paths: Sequence[str | Path]) -> List[str]:
    """Sheets present in every workbook (order of the first file)."""
    common: Optional[List[str]] = None
    for path in paths:
        names = list_sheet_names(path)
        common = names if common is None else [n for n in common if n in names]
    return common or []


def _pick_sheet(sheetnames: Sequence[str], requested: Optional[str]) -> tuple[str, Optional[str]]:
    """Requested sheet, then MASTER, then the first sheet. Returns (sheet, warning)."""
    if requested and requested in sheetnames:
        return requested, None
    upper = {name.upper(): name for name in sheetnames}
    chosen = upper.get("MASTER") or sheetnames[0]
    warning = f"Sheet '{requested}' not found; used '{chosen}'." if requested else None
    return chosen, warning


def parse_workbook(
    path: str | Path,
    *,
    sheet_name: Optional[str] = None,
    should_cancel: Optional[Callable[[], bool]] = None,
    progress: Optional[Callable[[int], None]] = None,
) -> ParsedWorkbook:
    """Parse one workbook into staged rows (grouping key: reconciled date)."""
    path = Path(path).expanduser().resolve()
    key = file_key_for(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if not wb.sheetnames:
            raise ValueError(f"{path.name}: workbook has no sheets.")
        used_sheet, sheet_warning = _pick_sheet(wb.sheetnames, sheet_name)
        loaded = LoadedFile(file_key=key, path=str(path), name=path.name, sheet=used_sheet)
        if sheet_warning:
            loaded.warnings.append(sheet_warning)
        result = ParsedWorkbook(file=loaded)

        header: Optional[List[Any]] = None
        cols: Dict[str, int] = {}
        workbook_date = parse_source_file_date(path.name)
        undated: List[StagedRow] = []

        for row_number, raw in enumerate(wb[used_sheet].iter_rows(values_only=True), start=1):
            if should_cancel is not None and row_number % 64 == 0 and should_cancel():
                raise DailyImportCancelled()
            values = list(raw or ())
            if header is None:
                if _is_blank_row(values):
                    continue
                mapped = _map_headers(values)
                if "description" in mapped and ({"tzs", "usd", "date"} & mapped.keys()):
                    header, cols = values, mapped
                elif _looks_like_classic_matumizi(values):
                    header, cols = values, dict(_CLASSIC_MATUMIZI_COLS)
                elif row_number >= _HEADER_SCAN_ROWS:
                    break
                continue
            if _is_blank_row(values):
                continue
            if progress is not None and row_number % 250 == 0:
                progress(row_number)
            _parse_data_row(
                result,
                header,
                cols,
                values,
                row_number,
                workbook_date=workbook_date,
                undated=undated,
            )

        if header is None:
            raise ValueError(
                f"{path.name}: no header row with Description and TZS/USD/Date "
                f"columns found in sheet '{used_sheet}'."
            )
        loaded.headers = _header_labels(header)
        if "source_file" not in cols and workbook_date is None:
            loaded.warnings.append(
                "No SOURCE FILE column and no date in the file name; "
                "using the most common row date as the reconciled date."
            )
        _assign_majority_date(result, undated)
        loaded.row_count = len(result.rows)
        loaded.skipped_count = len(result.skipped)
        return result
    finally:
        wb.close()


def _parse_data_row(
    result: ParsedWorkbook,
    header: Sequence[Any],
    cols: Dict[str, int],
    values: Sequence[Any],
    row_number: int,
    *,
    workbook_date: Optional[date],
    undated: List[StagedRow],
) -> None:
    loaded = result.file
    cells = _excel_cells(header, values)
    source_cell = " ".join(cell_str(_cell(values, cols, "source_file")).split())
    source_file = source_cell or loaded.name
    batch_date = parse_source_file_date(source_file) if source_cell else workbook_date
    if source_cell and batch_date is None:
        batch_date = workbook_date

    def skip(reason: str, detail: str = "") -> None:
        result.skipped.append(
            SkippedExcelRow(
                file_key=loaded.file_key,
                source_file=source_file,
                excel_row=row_number,
                reason=reason,
                detail=detail,
                batch_date=batch_date,
                excel_cells=cells,
            )
        )

    description = normalize_import_description(_cell(values, cols, "description"))
    if not description:
        skip(SKIP_EMPTY_DESCRIPTION)
        return
    if description.startswith("TOTAL") or cell_str(_cell(values, cols, "serial")).upper() == "TOTAL":
        skip(SKIP_TOTAL_ROW, description)
        return

    tzs_raw = _cell(values, cols, "tzs")
    usd_raw = _cell(values, cols, "usd")
    tzs = parse_amount(tzs_raw)
    usd = parse_amount(usd_raw)
    if tzs is None and usd is None:
        raw_text = cell_str(tzs_raw) or cell_str(usd_raw)
        if raw_text:
            skip(SKIP_UNPARSEABLE_AMOUNT, f"{description} | {raw_text}")
        else:
            skip(SKIP_MISSING_AMOUNT, description)
        return
    amount, amount_usd, currency = pack_money(tzs, usd)
    if not amount and not amount_usd:
        skip(SKIP_ZERO_AMOUNT, description)
        return

    if source_cell and batch_date is None:
        skip(SKIP_BAD_SOURCE_DATE, f"Cannot read a date from {source_cell!r}")
        return

    excel_dt = parse_date_value(_cell(values, cols, "date"))
    row = StagedRow(
        row_id=f"{loaded.file_key}|{loaded.sheet}|{row_number}",
        file_key=loaded.file_key,
        source_file=source_file,
        excel_row=row_number,
        batch_date=batch_date or date.min,
        description=description,
        amount=amount,
        currency=currency,
        amount_usd=amount_usd,
        excel_date=excel_dt.date() if excel_dt else None,
        truck_raw=" ".join(cell_str(_cell(values, cols, "truck")).upper().split()),
        lpo_do=cell_str(_cell(values, cols, "lpo")).upper(),
        do_number=cell_str(_cell(values, cols, "do")).upper(),
        memo=cell_str(_cell(values, cols, "memo")).upper(),
        notes=cell_str(_cell(values, cols, "notes")),
        receipt_status=cell_str(_cell(values, cols, "receipt")),
        ownership=cell_str(_cell(values, cols, "ownership")).upper(),
        approver=cell_str(_cell(values, cols, "approver")).upper(),
        serial=_parse_serial(_cell(values, cols, "serial")),
        excel_cells=cells,
    )
    if batch_date is None:
        undated.append(row)
    result.rows.append(row)


def _assign_majority_date(result: ParsedWorkbook, undated: List[StagedRow]) -> None:
    """Rows with no name-derived date take the file's most common row date."""
    if not undated:
        return
    counts = Counter(r.excel_date for r in undated if r.excel_date is not None)
    if not counts:
        pending = {id(r) for r in undated}
        result.rows = [r for r in result.rows if id(r) not in pending]
        for row in undated:
            result.skipped.append(
                SkippedExcelRow(
                    file_key=row.file_key,
                    source_file=row.source_file,
                    excel_row=row.excel_row,
                    reason=SKIP_BAD_SOURCE_DATE,
                    detail="No date in the file name and no row dates",
                    excel_cells=row.excel_cells,
                )
            )
        return
    top = max(counts.values())
    majority = min(d for d, n in counts.items() if n == top)
    for row in undated:
        row.batch_date = majority
    result.file.warnings.append(
        f"Reconciled date {majority.strftime('%d/%m/%Y')} taken from row dates."
    )
